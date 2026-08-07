"""
Tests for manuscript processor.
"""

from pathlib import Path

from docx import Document

from pipeline.manuscript import (
    ManuscriptProcessor,
)


def test_extract_references(
    tmp_path: Path,
):

    file_path = tmp_path / "article.docx"

    document = Document()

    document.add_paragraph("Introduction")

    document.add_paragraph("References")

    document.add_paragraph("Smith, J. (2020). Example article.")

    document.save(file_path)

    processor = ManuscriptProcessor()

    references = processor.extract_references(file_path)

    assert len(references) == 1

    assert references[0].year == 2020
