"""
End-to-end test for DOCX verification workflow.
"""

from pathlib import Path

from docx import Document

from models import (
    Provider,
    ReferenceMatch,
)
from report.exporter import ReportExporter
from report.generator import ReportGenerator
from verification.provider import VerificationProvider
from verification.service import VerificationService
from workflow.manuscript import ManuscriptProcessor
from workflow.verifier import ManuscriptVerifier


class MockProvider(VerificationProvider):
    """
    Deterministic provider for integration testing.
    """

    @property
    def name(self) -> str:
        return "mock"

    def search(self, reference):
        return [
            ReferenceMatch(
                provider=Provider.CROSSREF,
                title=reference.title,
                authors=reference.authors,
                year=reference.year,
                overall_score=1.0,
            )
        ]


def create_test_docx(path: Path) -> None:
    """
    Create minimal manuscript DOCX.
    """

    document = Document()

    document.add_paragraph("Introduction")

    document.add_paragraph("This is a test manuscript.")

    document.add_paragraph("References")

    document.add_paragraph("Smith, J. (2020). Example article.")

    document.save(path)


def test_complete_verification_workflow(tmp_path):
    """
    Test DOCX -> references -> verification -> report.
    """

    docx_path = tmp_path / "article.docx"

    create_test_docx(docx_path)

    processor = ManuscriptProcessor()

    references = processor.extract_references(docx_path)

    assert len(references) == 1

    service = VerificationService(providers=[MockProvider()])

    verifier = ManuscriptVerifier(service)

    results = verifier.verify_references(references)

    assert len(results) == 1

    report = ReportGenerator().generate_summary(results)

    assert report["total"] == 1

    output_dir = tmp_path / "output"

    exporter = ReportExporter()

    exporter.export_json(
        report,
        output_dir / "report.json",
    )

    exporter.export_csv(
        report,
        output_dir / "report.csv",
    )

    assert (output_dir / "report.json").exists()

    assert (output_dir / "report.csv").exists()
